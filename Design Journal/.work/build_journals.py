from __future__ import annotations

import shutil
from hashlib import sha256
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(r"C:\Users\yuyan\Desktop\S4\DECO 7285")
REFERENCE = ROOT / "Design Journal" / "Design_Journal_Format.docx"
OUTPUT_DIR = ROOT / "Design Journal" / "Completed"
EXPECTED_SHA = "7ec962d82172be98ec0d223367a6e69316611fdb36ed51d8fbf0860a8caaed4c"


ENTRIES = [
    {
        "week": "1",
        "slug": "problem_framing",
        "title": "Framing Friendship in a New Environment",
        "image": ROOT / "img" / "1.jpeg",
        "caption": "Early tutorial ideation mapping barriers to connection and possible design directions.",
        "description": [
            "Our project began with a broad question: how might we help people make friends when they enter a new place, life stage, hobby or community? During the first tutorial we mapped experiences around loneliness, unfamiliar environments and the difficulty of approaching people who already appear socially connected.",
            "The discussion showed that the problem was not simply a lack of events. A newcomer may know that activities exist but still feel uncertain about where they belong, how to enter a group and what to say. We therefore started to treat friendship as a gradual process rather than a single matching moment.",
        ],
        "learning": [
            "I learned that an initial design question can be useful without being a final problem definition. Our first question opened the space, but it was still too broad to guide a distinctive interaction. The most important next step was to gather evidence about the moments where connection breaks down, especially before and immediately after a first meeting.",
        ],
    },
    {
        "week": "2",
        "slug": "research_planning",
        "title": "Planning Contextual User Research",
        "image": ROOT / "img" / "2.jpeg",
        "caption": "Research planning activity used to identify participants, contexts and interview topics.",
        "description": [
            "We translated the broad topic into an interview plan focused on people living in shared accommodation and participating in local social or community groups. We prepared questions about where residents interact, what makes approaching others easier, what prevents connections from developing and what helps people stay in contact after meeting.",
            "We deliberately asked about observed behaviour and concrete situations rather than asking participants to design an app for us. Shared kitchens, lifts, common rooms, gyms and organised activities became important contexts because they create natural opportunities for repeated contact.",
        ],
        "learning": [
            "I learned that research questions should reveal behaviour, feelings and context before they test a preferred solution. Asking about real experiences helped us avoid assuming that technology alone would solve loneliness. It also gave us a clearer basis for recruiting participants and comparing patterns across interviews.",
        ],
    },
    {
        "week": "3",
        "slug": "interview_study",
        "title": "Interviewing Residents About Social Connection",
        "image": ROOT / "img" / "3.jpeg",
        "caption": "Interview question development focused on community spaces, barriers and continued contact.",
        "description": [
            "The interviews showed that new residents often begin with small, situational conversations: meeting a roommate, speaking in a kitchen or lift, asking about study, or attending an accommodation event. Shared interests and repeated encounters make interaction easier, while language confidence, established friendship groups, busy schedules and fear of awkwardness can prevent participation.",
            "A recurring insight was that quieter residents are not necessarily unwilling to socialise. They may need more time, a patient conversational partner and a shared activity that removes the pressure of introducing themselves without context. Maintaining contact also depended on a mutual reason to meet again rather than on exchanging contact details alone.",
        ],
        "learning": [
            "I learned to distinguish visible behaviour from motivation. Someone who appears quiet may be interested but uncertain, busy or communicating in a second language. This changed our design focus from pushing participation toward creating safer entry points, manageable group sizes and reasons for repeated interaction.",
        ],
    },
    {
        "week": "4",
        "slug": "insight_synthesis",
        "title": "Turning Interview Data into Design Principles",
        "image": ROOT / "img" / "4.jpeg",
        "caption": "Affinity mapping used to group interview findings into themes and design principles.",
        "description": [
            "We grouped the interview material into themes including difficulty entering an existing group, gradual motivation to connect, a spectrum of comfort, the role of particular spaces and common interests as a bridge. These themes led to practical principles: support smaller groups, provide a shared purpose, allow different levels of participation and help users continue a connection after the first meeting.",
            "This synthesis moved the project away from a generic event-discovery service. The emerging opportunity was to support the fragile transition between noticing an activity and feeling able to join it. We began to frame FriendZone as a bridge from online preparation to face-to-face participation.",
        ],
        "learning": [
            "I learned that synthesis is not just summarising comments. The value comes from connecting evidence to a design implication and being able to explain that connection. The affinity map made gaps visible and helped the team agree on principles that could later be tested through a prototype.",
        ],
    },
    {
        "week": "5",
        "slug": "storyboard_and_ethics",
        "title": "Storyboarding a Care-Centred Social Journey",
        "image": ROOT / "img" / "11.jpeg",
        "caption": "Early storyboard showing a newcomer moving from discovery to local connection.",
        "description": [
            "We created a storyboard around a newcomer entering Brisbane, selecting interests and joining a small local circle. The journey included discovery, a low-pressure message, an organised activity and mutual follow-up. This made the service understandable as a sequence rather than a collection of features.",
            "At the same time, we considered privacy, consent, safety, inclusion and the risk of turning friendship into measurable performance. The proposal therefore emphasised small circles, user choice and gradual participation. Rewards were treated as light encouragement rather than public competition, although we recognised that this assumption still required testing.",
        ],
        "learning": [
            "I learned that ethics changes the interaction itself. Protecting vulnerable users is not an extra checklist added after design; it affects what information is collected, how much control users have and whether prompts feel supportive or manipulative. Storyboarding helped reveal these ethical decisions at each stage of the journey.",
        ],
    },
    {
        "week": "6",
        "slug": "prototype_iterations",
        "title": "Building and Refining the FriendZone Prototype",
        "image": ROOT / "img" / "12.jpeg",
        "caption": "Refined storyboard for the care-centred FriendZone web application.",
        "description": [
            "The first interactive prototype introduced onboarding, interest selection, recommended circles, group chat, prompts and participation rewards. Subsequent iterations improved matching by recommending relevant circles, added optional profile details, expanded interest selection and introduced search, circle creation and an in-person mode.",
            "These iterations made the concept more complete, but they also exposed a risk: the application could become another broad social platform. We therefore reviewed which features directly supported the core transition from uncertainty to real-world participation and which features were only adding complexity.",
        ],
        "learning": [
            "I learned that adding functionality is not the same as strengthening a concept. Iteration should make the central interaction clearer, not simply make the prototype larger. The prototype became most useful when it allowed us to question our assumptions and identify which parts needed evidence or simplification.",
        ],
    },
    {
        "week": "7",
        "slug": "conversation_prompts",
        "title": "Focusing on Conversation Prompts",
        "image": ROOT / "img" / "13.jpeg",
        "caption": "Tutorial exploration of care, comfort and low-pressure interaction principles.",
        "description": [
            "We refined conversation prompts so that they could support users before and during a meeting without keeping attention on the screen. Prompts became more gradual as users spent time in a circle, moving from simple shared-interest questions toward more personal conversation only when familiarity had developed.",
            "This interaction responded directly to the interview finding that people often need a reason to speak. It also supported the idea of putting phones down once the conversation begins. The phone prepares the encounter, but the value of the experience comes from reciprocal, face-to-face interaction.",
        ],
        "learning": [
            "I learned that a small interaction can express the whole design position more clearly than a long feature list. Prompt timing, wording and the option to skip are important because a helpful suggestion can quickly become artificial or intrusive. The design must support conversation without scripting friendship.",
        ],
    },
    {
        "week": "8",
        "slug": "scope_after_critique",
        "title": "Narrowing the Concept After Studio Critique",
        "image": ROOT / "img" / "17.jpeg",
        "caption": "Studio board documenting the interaction journey, feedback and revised priorities.",
        "description": [
            "Studio feedback challenged us to explain why someone would choose FriendZone over Facebook events, Eventbrite or other meetup services. Small groups alone were not sufficiently distinctive because people can form smaller conversations inside a large event. We were advised to focus on one or two interactions, define a more specific audience and map research findings to requirements and features.",
            "The critique also questioned whether conventional points and rewards would be meaningful. It suggested exploring a tangible interaction and making face-to-face connection the centre of the design. We responded by pausing feature expansion and researching loneliness, phone use, rewards and the difference between social isolation and the feeling of loneliness.",
        ],
        "learning": [
            "I learned to treat critique as a request for sharper reasoning rather than more content. A narrow concept gives us a better chance to justify decisions and test them meaningfully. The strongest direction was not a complete social network, but an interaction that helps a specific user move into a real conversation.",
        ],
    },
    {
        "week": "9 to 10",
        "slug": "evaluation_and_redesign",
        "title": "Evaluating the Prototype and Changing Direction",
        "image": ROOT / "outputs" / "friendzone-evaluation-report" / "sus_scores.png",
        "caption": "System Usability Scale scores from four prototype participants, showing substantial variation.",
        "description": [
            "We evaluated the prototype with four anonymous participants using task-based questions and the System Usability Scale. The average SUS score was 63.1, with scores ranging from 20 to 90. Conversation prompts were consistently the strongest feature, while the NFC workflow, timer and leaderboard were confusing or unclear for some participants.",
            "Privacy and pressure were major concerns. Participants questioned what interaction data would be stored, who could see it and whether time together represented meaningful connection. Public ranking also felt competitive. Based on these findings and later research, we decided to retain the phone-down and prompt concepts while replacing the leaderboard with a friend relationship map that visualises connections without ranking people.",
        ],
        "learning": [
            "I learned that usability and social acceptability are different. A person may understand a feature and still reject its values or emotional effect. The next prototype should explain consent and data visibility, make prompts contextual and optional, and test a non-competitive relationship map on mobile devices with moderated observation.",
        ],
    },
]


def set_run_font(run, size=10.5, bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)


def write_text(paragraph, text, *, size=10.5, bold=False, italic=False, after=4, line=1.04):
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    return paragraph


def image_size(path: Path, max_w=5.9, max_h=2.15):
    with Image.open(path) as im:
        width, height = im.size
    ratio = min(max_w / width, max_h / height)
    return Inches(width * ratio), Inches(height * ratio)


def build_entry(entry):
    target = OUTPUT_DIR / f"Design_Journal_Week_{entry['week'].replace(' ', '_')}_{entry['slug']}.docx"
    shutil.copy2(REFERENCE, target)
    doc = Document(target)
    paragraphs = doc.paragraphs
    if len(paragraphs) < 26:
        raise RuntimeError("Unexpected template structure")

    write_text(
        paragraphs[0],
        "Student Name: Yuyang Gong      Student Number: __________      Team Name: FriendZone",
        size=9.4,
        after=2,
    )
    write_text(paragraphs[1], f"Week: {entry['week']}", size=10, after=3)

    clear_paragraph(paragraphs[2])
    label = paragraphs[2].add_run("Title: ")
    set_run_font(label, size=11, bold=True)
    value = paragraphs[2].add_run(entry["title"])
    set_run_font(value, size=11, bold=False)
    paragraphs[2].paragraph_format.space_after = Pt(5)
    paragraphs[2].paragraph_format.keep_with_next = True

    write_text(paragraphs[3], "", after=0)
    write_text(paragraphs[4], "Description:", size=10.5, bold=True, after=3)
    paragraphs[4].paragraph_format.keep_with_next = True
    for idx, text in enumerate(entry["description"]):
        write_text(paragraphs[5 + idx], text, size=10.0, after=4, line=1.02)

    visual_label_idx = 5 + len(entry["description"])
    write_text(paragraphs[visual_label_idx], "Visual Evidence:", size=10.5, bold=True, after=3)
    paragraphs[visual_label_idx].paragraph_format.keep_with_next = True

    image_p = paragraphs[visual_label_idx + 1]
    clear_paragraph(image_p)
    image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_p.paragraph_format.space_after = Pt(2)
    image_run = image_p.add_run()
    width, height = image_size(entry["image"])
    image_run.add_picture(str(entry["image"]), width=width, height=height)

    caption_p = paragraphs[visual_label_idx + 2]
    write_text(caption_p, entry["caption"], size=8.5, italic=True, after=5, line=1.0)
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    learning_label_idx = visual_label_idx + 3
    write_text(paragraphs[learning_label_idx], "Statement on Learning:", size=10.5, bold=True, after=3)
    paragraphs[learning_label_idx].paragraph_format.keep_with_next = True
    for idx, text in enumerate(entry["learning"]):
        write_text(paragraphs[learning_label_idx + 1 + idx], text, size=10.0, after=0, line=1.02)

    keep_until = learning_label_idx + len(entry["learning"])
    for paragraph in list(doc.paragraphs[keep_until + 1 :]):
        remove_paragraph(paragraph)

    doc.core_properties.title = entry["title"]
    doc.core_properties.subject = "DECO7285 Design Journal"
    doc.core_properties.author = "Yuyang Gong"
    doc.core_properties.keywords = "FriendZone, design journal, DECO7285"
    doc.save(target)
    return target


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    actual_sha = sha256(REFERENCE.read_bytes()).hexdigest()
    if actual_sha != EXPECTED_SHA:
        raise RuntimeError(f"Template changed: {actual_sha}")
    outputs = [build_entry(entry) for entry in ENTRIES]
    for output in outputs:
        print(output)


if __name__ == "__main__":
    main()
