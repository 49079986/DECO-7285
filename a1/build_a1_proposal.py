from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"C:\Users\yuyan\Desktop\S4\DECO 7285\a1\FriendZone_Project_Proposal.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_borders(cell, color="D9D9D9", size="8"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = margins.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            margins.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def remove_style_borders(style):
    p_pr = style._element.find(qn("w:pPr"))
    if p_pr is not None:
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is not None:
            p_pr.remove(p_bdr)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, size=10.4, bold=False, italic=False):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_para(doc, text="", style=None, before=0, after=4.4, line=1.04):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if text:
        r = p.add_run(text)
        set_font(r)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(5)
    for run in p.runs:
        set_font(run, size=14 if level == 1 else 12, bold=True)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.63)
section.bottom_margin = Inches(0.63)
section.left_margin = Inches(0.76)
section.right_margin = Inches(0.76)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10.4)
styles["Title"].font.name = "Arial"
styles["Title"].font.size = Pt(17)
styles["Title"].font.bold = True
remove_style_borders(styles["Title"])

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
remove_paragraph_borders(title)
r = title.add_run("FriendZone Project Proposal")
set_font(r, size=17, bold=True)

subtitle = add_para(doc, "A care-centred social web application for newcomers in Brisbane", after=12)
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in subtitle.runs:
    set_font(run, size=11, italic=True)

add_heading(doc, "Project Definition", 1)
add_para(doc, "FriendZone is a proposed social web application that helps newcomers in Brisbane build meaningful local connections through shared interests, gentle conversation prompts, and small daily rewards for showing up. The project responds to the question developed in tutorial work: how can we help people make new connections in a foreign environment? The initial problem space includes exchange students, people who have recently moved for work or study, stay-at-home partners, remote workers, and introverted users who may want social contact but feel anxious about entering unfamiliar groups.")
add_para(doc, "Moving to a new city often interrupts ordinary routines of friendship. Newcomers may not know where to meet people, which groups are suitable for beginners, how much social effort will be expected, or whether they will be forced into uncomfortable introductions. Existing social platforms often emphasise broad discovery, public profiles, fast messaging, or event attendance, but they do not always support the uncertain, gradual and emotionally sensitive process of becoming socially comfortable in a new place. FriendZone therefore treats connection as a paced journey rather than a quick match.")
add_para(doc, "The proposed setting is community-based social activity in Brisbane, especially small activities in public or semi-public places such as cafes, libraries, shared kitchens, rooftop lounges and walking routes. Users begin by entering a name for the prototype and choosing interests such as cooking, movies, study, gaming, board games, photography, running, hiking, coffee, music, basketball or yoga. The system then recommends small interest-based circles with visible member numbers, locations and recurring meeting days.")
add_para(doc, "The main user scenario follows Noah, a person who has recently moved to Brisbane and wants to meet people but feels nervous about approaching strangers. Noah selects photography, coffee and board games. FriendZone recommends circles such as Coffee Crawl, Library Buddies, Mt Coot-tha Walkers, Cooking Circle and Board Game Night. Each card shows a short description, current group size and meeting location. Noah joins Coffee Crawl, a small group that visits a new cafe on Sundays. Instead of expecting him to start a conversation from nothing, the chat offers a light prompt such as \"Flat white or long black?\" and makes participation feel manageable.")
add_para(doc, "The design solution has four main phases. First, onboarding makes interests visible without demanding a detailed public profile. Second, discovery recommends small circles rather than large anonymous events. Third, circle chat provides small prompts so that conversation can start around shared doing, not only self-presentation. Fourth, continuity is supported through daily participation, meetup RSVP, post-meetup bonus windows, streaks and badges. These rewards are intended as gentle encouragement for showing up, not as public competition or social ranking.")

add_heading(doc, "Significance of Proposed Interaction", 1)
add_para(doc, "The core interaction in FriendZone is a transition from digital preparation to embodied social participation. The application does not replace face-to-face connection with online chatting; instead, it supports people before, during and after local activities. This is important because the hardest part of making friends in a new environment is often the moment of entry: deciding whether a group feels safe, finding the meeting point, starting a first conversation, and knowing whether further contact is mutually welcome.")
add_para(doc, "The interaction qualities are care, predictability, agency, low pressure, reciprocity and continuity. Care is expressed through small circles, familiar interests and prompts that reduce the burden of inventing a topic. Predictability appears in group descriptions, member counts, locations and meeting schedules. Agency is central because users choose which circles to join and how much to contribute. Low pressure is achieved through the principle expressed in the prototype: something to do, not something to say. Reciprocity matters because a circle should support mutual participation rather than one-sided pursuit. Continuity is supported through a minimal daily ask: one message a day keeps the thread alive.")
add_para(doc, "The modalities combine form-based interest selection, visual circle cards, chat, reminder-like prompts, RSVP controls, points, badges and in-person social action. For example, a Coffee Crawl card shows that the group meets in West End on Sundays and has limited members. The circle chat displays messages from other members, offers a conversation prompt, and allows the user to send a short response. The prototype also includes demo controls that show how days, streaks and post-meetup bonuses could work over time.")
add_para(doc, "This interaction approach is significant because it recognises that social confidence varies. Some people are comfortable approaching a group immediately, while others need a path that starts with observation, a short message, or a shared activity. The design avoids assuming that extroverted behaviour is the norm. It also treats social participation as situated: users may behave differently depending on language, culture, past experience, time pressure, personal circumstances and the specific environment. In this sense, FriendZone aligns with values-oriented HCI, where values and ethics are not separate from design but should be surfaced as part of design practice (Shilton, 2018).")
add_para(doc, "The application also reflects the idea that technologies mediate human action. Verbeek (2006) argues that designed artefacts can shape what users do by inviting, discouraging or organising particular behaviours. FriendZone therefore must be designed carefully because its interface will influence how newcomers interpret social risk, how groups welcome participants, and how connections are continued. The aim is to inscribe gentle social support into the interaction while avoiding manipulative pressure.")

add_heading(doc, "Ethical Issues and Action Plan", 1)
add_para(doc, "The ethical issues in FriendZone are not secondary risks; they are central to whether the project succeeds. Because the application deals with social vulnerability, location, personal preferences and private relationships, the design must protect users while still helping them enter real community spaces. The following table identifies key ethical concerns and practical actions.")

table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [Inches(1.45), Inches(2.35), Inches(2.55), Inches(1.35)]
headers = ["Ethical issue", "Why it matters", "Action plan", "Design value"]
repeat_table_header(table.rows[0])
for i, header in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.width = widths[i]
    cell.text = header
    set_cell_shading(cell, "1F4E79")
    set_cell_borders(cell)
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        for run in p.runs:
            set_font(run, size=9, bold=True)
            run.font.color.rgb = RGBColor(255, 255, 255)

rows = [
    ("Privacy of personal preferences", "Interest choices, attendance and social activity can reveal sensitive information about identity and routines.", "Collect only necessary profile data; make privacy settings clear; allow users to edit or delete interests, joined circles and activity history.", "Privacy"),
    ("Consent in messaging", "A user may feel uncomfortable if joining a circle exposes them to unwanted direct pressure or contact.", "Keep early interaction group-based; provide mute, leave, block and report controls; make private messaging opt-in.", "Reciprocity"),
    ("Physical safety at activities", "The app moves users from online discovery into real-world meetings.", "Prefer public venues; verify hosts; show meeting points; provide check-in, emergency guidance and event reporting.", "Safety"),
    ("Pressure through rewards", "Points, streaks and badges could make friendship feel like a performance or obligation.", "Use private, low-stakes rewards; avoid leaderboards, streak loss and public comparison; frame points as encouragement for showing up.", "Care"),
    ("Exclusion and bias", "Matching by language or interests could isolate users or reproduce social bubbles.", "Offer diverse circle recommendations; include newcomer-friendly mixed groups; monitor whether some users are repeatedly excluded.", "Inclusion"),
    ("Host power and responsibility", "Hosts can see support needs and shape the tone of the group.", "Give hosts training prompts, privacy rules and moderation tools; limit what hosts can see; review complaints.", "Accountability"),
    ("Data retention", "Relationship histories and attendance records could become unnecessary long-term social data.", "Collect the minimum data needed; set retention periods; allow users to export and delete their data.", "Autonomy"),
    ("Cultural misunderstanding", "Newcomers may have different expectations around conversation, privacy and friendship.", "Use inclusive language, multilingual support, clear event norms and optional written activity instructions.", "Respect"),
]

for row_index, row in enumerate(rows, start=1):
    cells = table.add_row().cells
    for i, value in enumerate(row):
        cells[i].width = widths[i]
        cells[i].text = value
        set_cell_borders(cells[i])
        set_cell_margins(cells[i])
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if row_index % 2 == 0:
            set_cell_shading(cells[i], "EEF4FA")
        for p in cells[i].paragraphs:
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.02
            for run in p.runs:
                set_font(run, size=8.2)

add_para(doc, "These actions follow an anticipatory approach to ethics. Kudina and Verbeek (2019) show that technologies and values can change together as people encounter new forms of mediation. For FriendZone, this means privacy, comfort and social obligation should be evaluated through testing with newcomers and circle hosts, not assumed once at the start. Early prototype sessions should ask users what information they feel comfortable sharing, what types of prompts feel supportive, and when rewards or daily nudges become pressure. This feedback should be used to revise the interface before wider deployment.")

add_heading(doc, "Conclusion", 1)
add_para(doc, "FriendZone proposes a care-centred way to help newcomers form local relationships at their own pace. Its contribution is not simply to recommend events, but to redesign the fragile transition between wanting connection and entering a new social space. By combining interest-based circles, gentle prompts, small-group discovery, chat continuity, meetup participation and low-stakes rewards, the application supports people who may otherwise remain isolated. The ethical challenge is to ensure that the same technology that makes connection easier does not expose vulnerability, pressure users, or turn friendship into measurable performance. For this reason, privacy, consent, safety and inclusion must remain core design requirements throughout the project.")

add_heading(doc, "AI Acknowledgement", 1)
add_para(doc, "This proposal draft was prepared with assistance from ChatGPT/Codex to structure the report, refine academic wording, and organise ethical issues from the supplied tutorial materials, readings and FriendZone prototype. Before submission, the final author should review the content, revise it to reflect their own design decisions and add a shareable public link to the AI chat if required by the course.")

add_heading(doc, "References", 1)
refs = [
    "Kudina, O., & Verbeek, P.-P. (2019). Ethics from within: Google Glass, the Collingridge dilemma, and the mediated value of privacy. Science, Technology, & Human Values, 44(2), 291-314. https://doi.org/10.1177/0162243918793711",
    "Shilton, K. (2018). Values and ethics in human-computer interaction. Foundations and Trends in Human-Computer Interaction, 12(2), 107-171. https://doi.org/10.1561/1100000073",
    "Verbeek, P.-P. (2006). Materializing morality: Design ethics and technological mediation. Science, Technology, & Human Values, 31(3), 361-380. https://doi.org/10.1177/0162243905285847",
]
for ref in refs:
    p = add_para(doc, ref, after=5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.left_indent = Inches(0.25)

for p in doc.paragraphs:
    remove_paragraph_borders(p)
    for run in p.runs:
        if run.text:
            if not run.font.name:
                set_font(run)

doc.save(OUTPUT)
print(OUTPUT)
