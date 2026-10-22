import csv
import json
import statistics
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(r"C:\Users\yuyan\Desktop\download\evaluation_responses_rows.csv")
OUTPUT_DIR = Path(r"C:\Users\yuyan\Desktop\S4\DECO 7285\outputs\friendzone-evaluation-report")
OUTPUT = OUTPUT_DIR / "FriendZone_prototype_evaluation.docx"
CHART = OUTPUT_DIR / "sus_scores.png"


def set_run_font(run, size=10.5, bold=None, italic=None, color="1F2937"):
    run.font.name = "Arial"
    r_pr = run._element.get_or_add_rPr()
    r_pr.rFonts.set(qn("w:ascii"), "Arial")
    r_pr.rFonts.set(qn("w:hAnsi"), "Arial")
    r_pr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=105, start=120, bottom=105, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, bold=False, color="1F2937", align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(str(text if text is not None else ""))
    set_run_font(run, size=9.2, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def style_table(table, widths, header=True):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.width = Inches(widths[col_index])
            set_cell_margins(cell)
        if header and row_index == 0:
            repeat_header(row)
            for cell in row.cells:
                set_cell_shading(cell, "334155")
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in cell.paragraphs[0].runs:
                    set_run_font(run, size=9.2, bold=True, color="FFFFFF")
        elif row_index % 2 == 0:
            for cell in row.cells:
                set_cell_shading(cell, "F8FAFC")


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:pBdr"))
    if existing is not None:
        p_pr.remove(existing)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right", "between", "bar"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "nil")
        borders.append(border)
    p_pr.append(borders)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    set_run_font(run, size=8.5, color="6B7280")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def add_body(doc, text, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.16
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.12
    run = p.add_run(text)
    set_run_font(run)
    return p


def safe_json(value):
    try:
        return json.loads(value or "{}")
    except json.JSONDecodeError:
        return {}


OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
with SOURCE.open("r", encoding="utf-8-sig", newline="") as handle:
    raw_rows = list(csv.DictReader(handle))

responses = []
for index, row in enumerate(raw_rows, start=1):
    responses.append({
        "index": index,
        "row": row,
        "answers": safe_json(row.get("answers")),
        "sus": safe_json(row.get("sus_responses")),
    })

scores = [float(item["row"]["sus_score"]) for item in responses]
mean_score = statistics.mean(scores)
median_score = statistics.median(scores)

chart_width, chart_height = 1400, 650
image = Image.new("RGB", (chart_width, chart_height), "white")
draw = ImageDraw.Draw(image)
font_regular_path = r"C:\Windows\Fonts\arial.ttf"
font_bold_path = r"C:\Windows\Fonts\arialbd.ttf"
font_title = ImageFont.truetype(font_bold_path, 42)
font_axis = ImageFont.truetype(font_regular_path, 27)
font_label = ImageFont.truetype(font_bold_path, 28)
draw.text((90, 34), "SUS scores by participant", fill="#111827", font=font_title)
left, top, right, bottom = 120, 120, 1325, 540
for value in range(0, 101, 20):
    y = bottom - (value / 100) * (bottom - top)
    draw.line((left, y, right, y), fill="#E5E7EB", width=2)
    label = str(value)
    box = draw.textbbox((0, 0), label, font=font_axis)
    draw.text((left - 20 - (box[2] - box[0]), y - 14), label, fill="#4B5563", font=font_axis)
bar_colors = ["#4F6D8A", "#4F6D8A", "#4F6D8A", "#D97757"]
slot = (right - left) / len(scores)
bar_width = 135
for idx, (score, color) in enumerate(zip(scores, bar_colors)):
    center = left + slot * (idx + 0.5)
    x0, x1 = center - bar_width / 2, center + bar_width / 2
    y0 = bottom - (score / 100) * (bottom - top)
    draw.rectangle((x0, y0, x1, bottom), fill=color)
    score_text = f"{score:g}"
    score_box = draw.textbbox((0, 0), score_text, font=font_label)
    draw.text((center - (score_box[2] - score_box[0]) / 2, y0 - 40), score_text, fill="#111827", font=font_label)
    p_text = f"P{idx + 1}"
    p_box = draw.textbbox((0, 0), p_text, font=font_axis)
    draw.text((center - (p_box[2] - p_box[0]) / 2, bottom + 20), p_text, fill="#374151", font=font_axis)
mean_y = bottom - (mean_score / 100) * (bottom - top)
for x in range(left, right, 24):
    draw.line((x, mean_y, min(x + 13, right), mean_y), fill="#6B7280", width=3)
mean_text = f"Mean {mean_score:.1f}"
mean_box = draw.textbbox((0, 0), mean_text, font=font_axis)
draw.rectangle((right - (mean_box[2] - mean_box[0]) - 28, mean_y - 39, right, mean_y - 3), fill="white")
draw.text((right - (mean_box[2] - mean_box[0]) - 16, mean_y - 36), mean_text, fill="#4B5563", font=font_axis)
image.save(CHART)

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(31, 41, 55)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.16

title_style = doc.styles["Title"]
title_style.font.name = "Arial"
title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
title_style.font.size = Pt(22)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style.paragraph_format.space_after = Pt(5)

for style_name, size, before, after in (("Heading 1", 15, 12, 6), ("Heading 2", 12, 9, 4), ("Heading 3", 10.5, 7, 3)):
    style = doc.styles[style_name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

add_page_number(section.footer.paragraphs[0])

title = doc.add_paragraph(style="Title")
title.add_run("FriendZone Prototype Evaluation")
remove_paragraph_borders(title)
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(13)
set_run_font(subtitle.add_run("Usability Findings and Design Recommendations"), size=12, bold=True, color="475569")

add_body(
    doc,
    "This evaluation examined whether users could understand and complete the main FriendZone prototype tasks and how they perceived the conversation prompt, NFC interaction logging, leaderboard and continuity features. Four participants completed the task-based evaluation and the System Usability Scale. The results show an inconsistent usability experience: the mean SUS score was 63.1, but individual scores ranged from 20 to 90. Conversation prompts showed clear value, while the NFC workflow, leaderboard and data visibility require stronger explanation and user control.",
    space_after=9,
)

doc.add_paragraph("Evaluation Questions", style="Heading 1")
add_bullet(doc, "How usable is the FriendZone prototype when users select a group, use a conversation prompt and interpret the NFC timer and leaderboard?")
add_bullet(doc, "Which features support face to face interaction, and which features create confusion, privacy concerns or social pressure?")
add_bullet(doc, "What design changes should be prioritised before the next prototype iteration?")

doc.add_paragraph("Method", style="Heading 1")
doc.add_paragraph("Participants", style="Heading 2")
add_body(
    doc,
    "Four anonymous participants completed the evaluation. Two used the Chinese questionnaire and two used the English questionnaire. All participants used a desktop device, reported no prior use of FriendZone and gave permission for anonymous quotations. Because no demographic information was collected, the findings cannot be compared across age, gender or other user characteristics.",
)

doc.add_paragraph("Procedure", style="Heading 2")
add_body(
    doc,
    "Participants first reviewed the prototype and completed three tasks. Task 1 required selecting a relevant interest group. Task 2 examined the conversation prompt feature. Task 3 asked participants to interpret the NFC interaction, timer and leaderboard. After the tasks, participants completed the ten item System Usability Scale and answered open ended questions about helpful features, the leaderboard, privacy, continued use and design priorities.",
)
add_body(
    doc,
    "The current data set contains self reported task outcomes and written comments, but it does not contain a verified audio or video think aloud record. The present report therefore treats the qualitative data as open ended responses rather than claiming think aloud findings. A moderated think aloud guide is included in the appendix for the next evaluation round.",
)

doc.add_paragraph("Measures and Analysis", style="Heading 2")
add_body(
    doc,
    "SUS was selected because it provides a short global measure of perceived usability and has been evaluated across a wide range of interactive products (Bangor, Kortum, & Miller, 2008). For each participant, one was subtracted from each odd numbered response, each even numbered response was subtracted from five, and the sum was multiplied by 2.5. Written responses were reviewed using a light thematic analysis informed by Braun and Clarke (2006). Repeated comments were grouped into four themes: clarity and guidance, conversation support, competition and pressure, and privacy and data control.",
)

doc.add_paragraph("Results", style="Heading 1")
doc.add_paragraph("System Usability Scale", style="Heading 2")
add_body(
    doc,
    f"The mean SUS score was {mean_score:.1f}, the median was {median_score:.1f}, and scores ranged from {min(scores):g} to {max(scores):g}. The wide range is more informative than the mean alone. Participants 1 and 2 rated the prototype highly, while Participant 4 gave a score of 20 and reported difficulty understanding several functions. With only four participants, the mean should be treated as an exploratory indicator rather than a representative estimate of the target population.",
)

sus_table = doc.add_table(rows=1, cols=5)
for col, text in enumerate(["Participant", "Language", "Task 1", "Task 3", "SUS score"]):
    set_cell_text(sus_table.rows[0].cells[col], text, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
for item in responses:
    answers = item["answers"]
    row = item["row"]
    values = [
        f"P{item['index']}",
        row.get("language", ""),
        answers.get("task_1_status") or "No status",
        answers.get("task_3_status") or "No status",
        row.get("sus_score", ""),
    ]
    cells = sus_table.add_row().cells
    for col, value in enumerate(values):
        set_cell_text(cells[col], value, align=WD_ALIGN_PARAGRAPH.CENTER)
style_table(sus_table, widths=[1.15, 1.0, 1.45, 1.45, 1.1])

chart_p = doc.add_paragraph()
chart_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
chart_p.paragraph_format.space_before = Pt(7)
chart_p.paragraph_format.space_after = Pt(2)
chart_p.add_run().add_picture(str(CHART), width=Inches(6.25))
caption = doc.add_paragraph()
caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
caption.paragraph_format.space_after = Pt(7)
set_run_font(caption.add_run("Figure 1  Individual SUS scores and the study mean"), size=9, italic=True, color="475569")

doc.add_paragraph("Task Findings", style="Heading 2")
add_body(
    doc,
    "Three participants completed the group selection task independently, while one marked the task as not completed. All four participants reported completing the conversation prompt task independently. Three participants reported independent completion of the NFC, timer and leaderboard task. The fourth participant did not provide a completion status and wrote that these functions were not clear. This pattern indicates that the main usability risk lies in explaining the transition from phone contact to recorded interaction and leaderboard output.",
)

doc.add_paragraph("Qualitative Themes", style="Heading 2")
doc.add_paragraph("Clarity and Guidance", style="Heading 3")
add_body(
    doc,
    "The lowest scoring participant could not clearly explain the NFC, start and end of the timer, or leaderboard. Other participants could describe the intended workflow, but their explanations focused on what they believed the system recorded rather than on visible interface feedback. This suggests that the prototype needs clearer system status, confirmation and explanation at the point where phones are placed together.",
)

doc.add_paragraph("Conversation Support", style="Heading 3")
add_body(
    doc,
    "Two participants explicitly identified the conversation prompt as helpful. One wrote that it could help quieter students integrate into a group, and another said that it helps when people do not know what to say. However, one participant found a suggested prompt artificial, and another received a response that did not make sense. The feature should therefore remain in the design, but prompts need stronger contextual relevance and an easy option to replace or skip them.",
)

doc.add_paragraph("Competition and Pressure", style="Heading 3")
add_body(
    doc,
    "Responses to the leaderboard and continuity features were mixed. One participant described the leaderboard as unnecessary and potentially competitive. Another saw some motivational value but worried that group members could infer a person's itinerary. Two participants described continuity or streak features as pressure. This result is consistent with experimental gamification research showing that leaderboards can increase task output without necessarily increasing intrinsic motivation (Mekler et al., 2017). A public individual ranking may therefore conflict with FriendZone's goal of comfortable and voluntary social connection.",
)

doc.add_paragraph("Privacy and Data Control", style="Heading 3")
add_body(
    doc,
    "All four responses raised either privacy, surveillance or control concerns. Participants questioned what information would appear after phone contact, whether group members could infer their activities and whether the recorded time represented real interaction. Research on location disclosure similarly shows that willingness to share depends on who can see the data, why it is requested and how much detail is disclosed (Consolvo et al., 2005). The interface should explain what is stored, who can see it and how a user can edit, delete or decline the record.",
)

doc.add_paragraph("Discussion", style="Heading 1")
add_body(
    doc,
    "FriendZone's phone down concept is supported by research showing that the presence of mobile devices can reduce perceived conversation quality, trust, empathy and closeness during face to face interaction (Przybylski & Weinstein, 2013; Misra et al., 2016). The prototype should therefore minimise phone interaction once a meeting begins. NFC can function as a brief transition into face to face contact, but it should not require repeated checking or create uncertainty about monitoring.",
)
add_body(
    doc,
    "The findings also show that usability and social acceptability cannot be treated as the same issue. A user may understand how the leaderboard works but still reject it because it feels competitive or intrusive. The next iteration should therefore evaluate both task success and the emotional effect of ranking, tracking and streak mechanisms.",
)

doc.add_paragraph("Design Recommendations", style="Heading 1")
recommendations = [
    ("NFC workflow is unclear", "One participant could not explain NFC, timer or leaderboard", "Add a short pre contact explanation, connection feedback and a confirmation screen before saving", "High"),
    ("Interaction records feel intrusive", "Participants mentioned surveillance, itinerary inference and inaccurate recorded time", "Show exactly what is stored and who can view it; allow opt out, edit and delete controls", "High"),
    ("Public ranking creates pressure", "Several responses associated leaderboards or streaks with pressure or comparison", "Replace default individual ranking with optional personal progress or a shared group goal", "High"),
    ("Prompts are useful but inconsistent", "Two participants found prompts helpful, while others found wording artificial or irrelevant", "Keep the feature but improve contextual prompts and provide skip, refresh and custom question options", "Medium"),
]
rec_table = doc.add_table(rows=1, cols=4)
for col, text in enumerate(["Issue", "Evidence", "Recommendation", "Priority"]):
    set_cell_text(rec_table.rows[0].cells[col], text, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
for issue, evidence, recommendation, priority in recommendations:
    cells = rec_table.add_row().cells
    for col, value in enumerate([issue, evidence, recommendation, priority]):
        set_cell_text(cells[col], value, align=WD_ALIGN_PARAGRAPH.CENTER if col == 3 else WD_ALIGN_PARAGRAPH.LEFT)
style_table(rec_table, widths=[1.25, 2.05, 2.55, 0.65])

doc.add_paragraph("Limitations", style="Heading 1")
add_bullet(doc, "The sample contained only four participants, so the SUS mean and qualitative patterns cannot be generalised to the wider target population.")
add_bullet(doc, "The data were self reported. The current export does not contain observation notes, completion times, recordings or verified think aloud data.")
add_bullet(doc, "Participants used a prototype rather than a fully implemented NFC system, so privacy and tracking responses were partly based on their interpretation of the concept.")
add_bullet(doc, "All participants used a desktop device even though the proposed interaction depends on mobile phones. The next test should use mobile devices and a functional or simulated NFC flow.")
add_bullet(doc, "Small usability samples can reveal important problems, but the number of problems discovered depends on how likely each problem is to occur (Lewis, 1994).")

doc.add_paragraph("Conclusion", style="Heading 1")
add_body(
    doc,
    "The evaluation identified both a useful social support mechanism and a major design risk. Conversation prompts can help users begin face to face interaction, but the NFC workflow, leaderboard and continuity mechanisms may create confusion, privacy concerns and social pressure. The next prototype should preserve the prompt feature, simplify the phone contact flow, make all recorded data transparent and replace default public ranking with a less competitive form of progress. A moderated mobile think aloud session should then test whether users understand these changes without assistance.",
)

doc.add_page_break()
doc.add_paragraph("References", style="Heading 1")
references = [
    "Bangor, A., Kortum, P. T., & Miller, J. T. (2008). An empirical evaluation of the System Usability Scale. International Journal of Human Computer Interaction, 24(6), 574-594. https://doi.org/10.1080/10447310802205776",
    "Braun, V., & Clarke, V. (2006). Using thematic analysis in psychology. Qualitative Research in Psychology, 3(2), 77-101. https://doi.org/10.1191/1478088706qp063oa",
    "Consolvo, S., Smith, I. E., Matthews, T., LaMarca, A., Tabert, J., & Powledge, P. (2005). Location disclosure to social relations: Why, when, and what people want to share. Proceedings of the SIGCHI Conference on Human Factors in Computing Systems, 81-90. https://doi.org/10.1145/1054972.1054985",
    "Lewis, J. R. (1994). Sample sizes for usability studies: Additional considerations. Human Factors, 36(2), 368-378. https://doi.org/10.1177/001872089403600215",
    "Mekler, E. D., Bruehlmann, F., Tuch, A. N., & Opwis, K. (2017). Towards understanding the effects of individual gamification elements on intrinsic motivation and performance. Computers in Human Behavior, 71, 525-534. https://doi.org/10.1016/j.chb.2015.08.048",
    "Misra, S., Cheng, L., Genevie, J., & Yuan, M. (2016). The iPhone effect: The quality of in person social interactions in the presence of mobile devices. Environment and Behavior, 48(2), 275-298. https://doi.org/10.1177/0013916514539755",
    "Przybylski, A. K., & Weinstein, N. (2013). Can you connect with me now? How the presence of mobile communication technology influences face to face conversation quality. Journal of Social and Personal Relationships, 30(3), 237-246. https://doi.org/10.1177/0265407512453827",
    "van den Haak, M. J., de Jong, M. D. T., & Schellens, P. J. (2004). Employing think aloud protocols and constructive interaction to test the usability of online library catalogues: A methodological comparison. Interacting with Computers, 16(6), 1153-1170. https://doi.org/10.1016/j.intcom.2004.07.007",
]
for reference in references:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.28)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.08
    set_run_font(p.add_run(reference), size=9.7)

doc.add_page_break()
doc.add_paragraph("Appendix Think Aloud Session Guide", style="Heading 1")
add_body(
    doc,
    "Use this guide for one moderated mobile session. The moderator should explain that the prototype is being tested, not the participant. Ask the participant to say what they are looking at, what they expect to happen and why they choose each action. Avoid explaining the interface until the participant becomes unable to continue.",
)
doc.add_paragraph("Moderator Opening", style="Heading 2")
add_body(
    doc,
    "Thank you for taking part. We are testing the FriendZone prototype, not your ability. As you complete each task, please say what you are thinking, what you expect each control to do and anything that feels confusing. I may remind you to keep talking, but I will not explain the interface unless you cannot continue.",
)
doc.add_paragraph("Tasks", style="Heading 2")
add_bullet(doc, "Task 1  Choose a group that matches your interests and explain what information influenced your choice.")
add_bullet(doc, "Task 2  Use a conversation prompt to begin an interaction and explain whether the prompt feels natural and relevant.")
add_bullet(doc, "Task 3  Explain what you expect to happen when two phones use NFC, when timing begins and ends, and what information appears on the leaderboard.")
doc.add_paragraph("Neutral Prompts", style="Heading 2")
add_bullet(doc, "What are you looking for now?")
add_bullet(doc, "What do you expect will happen next?")
add_bullet(doc, "What does this information mean to you?")
add_bullet(doc, "What, if anything, feels unclear?")
doc.add_paragraph("Observation Checklist", style="Heading 2")
check_table = doc.add_table(rows=1, cols=3)
for col, text in enumerate(["Observation", "What to record", "Evidence"]):
    set_cell_text(check_table.rows[0].cells[col], text, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
check_items = [
    ("Task success", "Completed independently, with prompting, or not completed", "Final screen and participant action"),
    ("Navigation", "Backtracking, hesitation or incorrect selection", "Screen or step where it occurred"),
    ("Expectation", "What the participant predicts before selecting a control", "Short quotation"),
    ("NFC understanding", "What the participant believes is recorded and shared", "Explanation in their own words"),
    ("Emotional response", "Comfort, pressure, privacy concern or confidence", "Observed reaction and quotation"),
]
for values in check_items:
    cells = check_table.add_row().cells
    for col, value in enumerate(values):
        set_cell_text(cells[col], value)
style_table(check_table, widths=[1.35, 2.65, 2.65])

for paragraph in doc.paragraphs:
    paragraph.paragraph_format.widow_control = True

doc.core_properties.title = "FriendZone Prototype Evaluation"
doc.core_properties.subject = "Usability findings and design recommendations"
doc.core_properties.author = "FriendZone Evaluation Team"
doc.save(OUTPUT)
print(f"OUTPUT={OUTPUT}")
print(f"RESPONSES={len(responses)}")
print(f"MEAN_SUS={mean_score:.1f}")
