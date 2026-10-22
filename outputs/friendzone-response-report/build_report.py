import csv
import json
import statistics
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(r"C:\Users\yuyan\Desktop\download\evaluation_responses_rows.csv")
OUTPUT_DIR = Path(r"C:\Users\yuyan\Desktop\S4\DECO 7285\outputs\friendzone-response-report")
OUTPUT = OUTPUT_DIR / "FriendZone_evaluation_responses.docx"


def set_run_font(run, western="Arial", east_asia="Microsoft YaHei", size=None, bold=None, italic=None, color=None):
    run.font.name = western
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run._element.rPr.rFonts.set(qn("w:ascii"), western)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), western)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cell_text(cell, text, bold=False, color="1F2937", align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(str(text if text is not None else ""))
    set_run_font(run, size=9.5, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def style_table(table, header=True, widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            set_cell_margins(cell)
        if header and row_idx == 0:
            set_repeat_table_header(row)
            for cell in row.cells:
                set_cell_shading(cell, "334155")
                for run in cell.paragraphs[0].runs:
                    set_run_font(run, size=9.5, bold=True, color="FFFFFF")
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif row_idx % 2 == 0:
            for cell in row.cells:
                set_cell_shading(cell, "F8FAFC")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    set_run_font(run, size=8.5, color="6B7280")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:pBdr"))
    if existing is not None:
        p_pr.remove(existing)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right", "between", "bar"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "nil")
        borders.append(border)
    p_pr.append(borders)


def add_question_answer(doc, question, answer):
    q = doc.add_paragraph()
    q.paragraph_format.space_before = Pt(7)
    q.paragraph_format.space_after = Pt(2)
    q.paragraph_format.keep_with_next = True
    qr = q.add_run(question)
    set_run_font(qr, size=10.5, bold=True, color="1F2937")
    a = doc.add_paragraph()
    a.paragraph_format.space_after = Pt(5)
    a.paragraph_format.line_spacing = 1.15
    ar = a.add_run(str(answer if answer not in (None, "") else "No response"))
    set_run_font(ar, size=10.5, color="374151")


def add_task_section(doc, number, title, status, response_label, response, difficulty):
    heading = doc.add_paragraph(style="Heading 2")
    heading.paragraph_format.keep_with_next = True
    heading.add_run(f"Task {number} {title}")
    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(4)
    label = meta.add_run("Completion status  ")
    set_run_font(label, size=9.5, bold=True, color="475569")
    value = meta.add_run(str(status or "No response"))
    set_run_font(value, size=9.5, color="475569")
    add_question_answer(doc, response_label, response)
    add_question_answer(doc, "Reported difficulty", difficulty)


with SOURCE.open("r", encoding="utf-8-sig", newline="") as handle:
    reader = csv.DictReader(handle)
    raw_rows = list(reader)

responses = []
for index, row in enumerate(raw_rows, start=1):
    answers = json.loads(row.get("answers") or "{}")
    sus = json.loads(row.get("sus_responses") or "{}")
    responses.append({"index": index, "row": row, "answers": answers, "sus": sus})

scores = [float(item["row"]["sus_score"]) for item in responses if item["row"].get("sus_score")]
average_score = statistics.mean(scores) if scores else 0
median_score = statistics.median(scores) if scores else 0

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(31, 41, 55)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.15

title_style = styles["Title"]
title_style.font.name = "Arial"
title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
title_style.font.size = Pt(22)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style.paragraph_format.space_after = Pt(7)

for style_name, size, before, after in (("Heading 1", 15, 12, 6), ("Heading 2", 12, 9, 4)):
    style = styles[style_name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

footer = section.footer
add_page_number(footer.paragraphs[0])

title = doc.add_paragraph(style="Title")
title.add_run("FriendZone Evaluation Responses")
remove_paragraph_borders(title)
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(14)
run = subtitle.add_run("SUS and Interview Findings")
set_run_font(run, size=12, bold=True, color="475569")

intro = doc.add_paragraph()
intro.add_run(
    f"This report organises {len(responses)} anonymous FriendZone evaluation responses into a readable format. "
    "It combines SUS scores, task feedback and interview comments while preserving each participant's original wording. "
    f"Scores range from {min(scores):g} to {max(scores):g}, which shows that participants had substantially different usability experiences."
)

doc.add_paragraph("Response Overview", style="Heading 1")
summary_table = doc.add_table(rows=1, cols=5)
headers = ["Participant", "Language", "Device", "Prior use", "SUS score"]
for col, text in enumerate(headers):
    set_cell_text(summary_table.rows[0].cells[col], text, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
for item in responses:
    row = item["row"]
    cells = summary_table.add_row().cells
    values = [
        f"Participant {item['index']}",
        row.get("language", ""),
        row.get("device", ""),
        row.get("prior_use", ""),
        row.get("sus_score", ""),
    ]
    for col, value in enumerate(values):
        set_cell_text(cells[col], value, align=WD_ALIGN_PARAGRAPH.CENTER)
style_table(summary_table, widths=[1.45, 1.1, 1.2, 1.2, 1.0])

stats = doc.add_paragraph()
stats.paragraph_format.space_before = Pt(8)
lead = stats.add_run("SUS summary  ")
set_run_font(lead, size=10.5, bold=True)
text = stats.add_run(f"Average {average_score:.1f}  Median {median_score:.1f}  Range {min(scores):g} to {max(scores):g}")
set_run_font(text, size=10.5)

doc.add_paragraph("Patterns Across Responses", style="Heading 1")
patterns = [
    "Usability was inconsistent. Two participants scored the interface at 80 or above, while one participant scored it at 20 and reported that several functions were unclear.",
    "The conversation prompt was explicitly identified as helpful by two participants because it could help quieter users start a discussion.",
    "The leaderboard received mixed reactions. Participants mentioned comparison, pressure and the possibility of inferring another person's activity from the shared record.",
    "NFC and activity tracking raised privacy and clarity concerns. One participant considered the approach less invasive than real-time location tracking, while others worried about unwanted information sharing or feeling watched.",
]
for item in patterns:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.add_run(item)

for item in responses:
    row = item["row"]
    answers = item["answers"]
    sus = item["sus"]
    doc.add_page_break()

    doc.add_paragraph(f"Participant {item['index']} {row.get('response_code', '')}", style="Heading 1")

    meta_table = doc.add_table(rows=5, cols=2)
    meta_values = [
        ("Submitted", row.get("submitted_at", "")),
        ("Language", row.get("language", "")),
        ("Device", row.get("device", "")),
        ("Prior use", row.get("prior_use", "")),
        ("Quote permission", row.get("quote_permission", "")),
    ]
    for ridx, (label, value) in enumerate(meta_values):
        set_cell_text(meta_table.rows[ridx].cells[0], label, bold=True)
        set_cell_text(meta_table.rows[ridx].cells[1], value)
        set_cell_shading(meta_table.rows[ridx].cells[0], "E2E8F0")
    style_table(meta_table, header=False, widths=[1.65, 4.6])

    doc.add_paragraph("SUS Results", style="Heading 2")
    score_p = doc.add_paragraph()
    score_p.paragraph_format.space_after = Pt(3)
    score_lead = score_p.add_run("Overall score  ")
    set_run_font(score_lead, size=10.5, bold=True)
    score_run = score_p.add_run(str(row.get("sus_score", "")))
    set_run_font(score_run, size=10.5)
    ratings = [answers.get(f"sus_{n}", sus.get(f"q{n}", "")) for n in range(1, 11)]
    ratings_p = doc.add_paragraph()
    ratings_p.paragraph_format.space_after = Pt(7)
    ratings_lead = ratings_p.add_run("Item ratings 1 to 10  ")
    set_run_font(ratings_lead, size=10.5, bold=True)
    ratings_run = ratings_p.add_run(", ".join(str(value) for value in ratings))
    set_run_font(ratings_run, size=10.5)

    doc.add_paragraph("Task Feedback", style="Heading 1")
    add_task_section(
        doc,
        1,
        "Group Selection",
        answers.get("task_1_status"),
        "Selected group and reason",
        answers.get("task_1_choice"),
        answers.get("task_1_difficulty"),
    )
    add_task_section(
        doc,
        2,
        "Conversation Prompt",
        answers.get("task_2_status"),
        "Prompt experience",
        answers.get("task_2_choice"),
        answers.get("task_2_difficulty"),
    )
    add_task_section(
        doc,
        3,
        "NFC Timer and Leaderboard",
        answers.get("task_3_status"),
        "Interpretation of the feature",
        answers.get("task_3_choice"),
        answers.get("task_3_difficulty"),
    )

    doc.add_paragraph("Interview Responses", style="Heading 1")
    add_question_answer(doc, "Which feature was most helpful", answers.get("most_helpful"))
    add_question_answer(doc, "How did the leaderboard make you feel", answers.get("board_reaction"))
    add_question_answer(doc, "What did you think about NFC and privacy", answers.get("nfc_privacy"))
    add_question_answer(doc, "Would the continuity or streak feature motivate continued use", answers.get("continuity"))
    add_question_answer(doc, "What would you change or prioritise", answers.get("priority_change"))

for paragraph in doc.paragraphs:
    paragraph.paragraph_format.widow_control = True

doc.core_properties.title = "FriendZone Evaluation Responses"
doc.core_properties.subject = "SUS scores task feedback and interview responses"
doc.core_properties.author = "FriendZone Evaluation Team"

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(f"OUTPUT={OUTPUT}")
print(f"RESPONSES={len(responses)}")
print(f"AVERAGE_SUS={average_score:.1f}")
